"""
Núcleo do Importador de Tomadores de Serviço — Domínio Folha (FOSERVICOS / FOVIGENCIAS_SERVICO).

Funciona sem acesso ao banco:
- referencia/ (gerada por gerar_referencia.py): CNAE 2.3, relação CNAE x FPAS, terceiros, RAT por CNAE,
  municípios, colunas e chaves das tabelas; leiaute/: leiaute da importação.
- Consultas públicas de CNPJ (BrasilAPI -> CNPJá -> ReceitaWS, com limite de vazão global e cache) e CEP (ViaCEP).
- A empresa prestadora define o perfil tributário (FPAS, terceiros, RAT, INSS, SEFIP, GPS); todo tomador herda
  esse perfil, venha da digitação ou da planilha (Soluções 136 e 3665 da Central Domínio).
- Saída: TXT da importação de tabelas (tab, cp1252, sem cabeçalho) + leiaute ajustado sem os campos que ficam nulos.
"""
import csv
import datetime as dt
import hashlib
import io
import json
import os
import re
import threading
import time
import unicodedata
import urllib.error
import urllib.request
import zipfile
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

T_SERV, T_VIG = "FOSERVICOS", "FOVIGENCIAS_SERVICO"
TABELAS = (T_SERV, T_VIG)
TIPO_TOMADOR = 4
SEFIP_TOMADOR = 150                           # Solução 136
FIM_ABERTO = dt.date(3000, 12, 31)
CORTE_CNAE_ESOCIAL = dt.date(2018, 12, 1)     # CNAE eSocial só para vigências anteriores a 12/2018 (Solução 3665)
ORIGEM_IMP_TABELAS = 1                        # origem_reg: 1 = Imp. Tabelas
RAT_INFORMADO = 1                             # TIPO_INFORMACAO_ALIQUOTA_ACIDENTE_TRABALHO: 1 = Informado

INTEIROS = {"integer", "smallint", "tinyint", "bigint", "unsigned int", "unsigned smallint", "unsigned bigint", "bit"}
LIMITES_INT = {"tinyint": (0, 255), "smallint": (-32768, 32767), "integer": (-2**31, 2**31 - 1), "bit": (0, 1)}
DECIMAIS = {"numeric", "decimal"}
TEXTOS = {"char", "varchar", "long varchar", "nchar", "nvarchar"}
DATAS = {"date", "timestamp", "datetime", "smalldatetime"}

# Regime da prestadora -> encargos do serviço (Solução 3665) e GPS; SEFIP do tomador = 150 (Solução 136)
REGIMES = {
    "LRLP": {"rotulo": "Lucro Real / Lucro Presumido", "terceiros": "sem_convenio", "inss": 20, "prolab": 20,
             "gps": 2100},
    "SN": {"rotulo": "Simples Nacional — Anexos I a III e V", "terceiros": "zero", "inss": 0, "prolab": 0,
           "gps": 2003},
    "SN4": {"rotulo": "Simples Nacional — Anexo IV", "terceiros": "zero", "inss": 20, "prolab": 20, "gps": 2100},
    "SNC": {"rotulo": "Simples Nacional — Concomitante", "terceiros": "zero", "inss": 20, "prolab": 20,
            "gps": 2100},
    "MEI": {"rotulo": "MEI", "terceiros": "zero", "inss": 3, "prolab": 0, "gps": 2100},
}
# Atividades do Anexo IV (LC 123, art. 18, § 5º-C): construção, vigilância, limpeza/conservação, advocacia
ANEXO_IV_PREFIXOS = ("41", "42", "43", "8011", "8111", "8121", "8129", "6911")

TIPOS_SERVICO_TOMADOR = {
    1: "1 - Cessão de mão de obra (inclusive contratante de cooperativa PJ)",
    2: "2 - Trabalho temporário - substituição de pessoal",
    3: "3 - Trabalho temporário - serviços extraordinários",
}
LOTACOES = {  # Tabela 10 do eSocial (S-1.1); confirme o código gravado pela Domínio num tomador cadastrado pela tela
    3: "03 - PF tomadora de cessão de mão de obra",
    4: "04 - PJ tomadora de cessão de mão de obra",
    5: "05 - PJ tomadora de cooperados (cooperativa de trabalho)",
    6: "06 - Entidade beneficente/isenta tomadora de cooperados",
    7: "07 - PF tomadora de cooperados",
    8: "08 - Operador portuário (avulsos)",
    9: "09 - Contratante de avulsos não portuários (sindicato)",
}
TIPOS_LOGRADOURO = {"RUA", "AVENIDA", "TRAVESSA", "ALAMEDA", "PRACA", "RODOVIA", "ESTRADA", "LARGO", "VIA", "QUADRA",
                    "VIELA", "BECO", "PARQUE", "SETOR", "VILA", "LADEIRA", "PASSAGEM", "CAMINHO", "AV", "R"}


# =====================================================================
# Texto, números e documentos
# =====================================================================
def norm(txt_) -> str:
    if txt_ is None or (isinstance(txt_, float) and pd.isna(txt_)):
        return ""
    s = unicodedata.normalize("NFKD", str(txt_))
    s = "".join(c for c in s if not unicodedata.combining(c)).upper()
    s = re.sub(r"[^A-Z0-9%]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def txt(v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    s = str(v).replace("\u00a0", " ").replace("\u200b", "").strip()
    return "" if s.lower() in ("nan", "none", "nat") else s


def digitos(v) -> str:
    return re.sub(r"\D", "", txt(v))


def _numero_excel(s: str) -> str:
    """'191.0' -> '191'; '3.3000167000101e+13' -> '33000167000101' (repr do Python, sem perda)."""
    if re.fullmatch(r"\d+\.0+", s):
        return s.split(".")[0]
    if re.fullmatch(r"\d+(\.\d+)?[eE]\+?\d+", s):
        try:
            d = Decimal(s)
            if d == d.to_integral_value():
                return str(int(d))
        except InvalidOperation:
            pass
    return s


def _dv_mod11(base: str, pesos) -> str:
    s = sum((ord(c) - 48) * p for c, p in zip(base, pesos))
    r = s % 11
    return "0" if r < 2 else str(11 - r)


def cnpj_valido(c: str) -> bool:
    """CNPJ numérico ou alfanumérico (IN RFB 2.229/2024): valor do caractere = código ASCII - 48."""
    if not re.fullmatch(r"[0-9A-Z]{12}\d{2}", c or "") or len(set(c)) == 1:
        return False
    p1 = [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]
    return c[12] == _dv_mod11(c[:12], p1) and c[13] == _dv_mod11(c[:13], [6] + p1)


def cpf_valido(c: str) -> bool:
    if not re.fullmatch(r"\d{11}", c or "") or len(set(c)) == 1:
        return False
    return c[9] == _dv_mod11(c[:9], range(10, 1, -1)) and c[10] == _dv_mod11(c[:10], range(11, 1, -1))


def normalizar_inscricao(v):
    """Valor digitado ou da planilha -> (inscrição limpa, aviso, erro)."""
    s = txt(v)
    if not s:
        return "", "", ""
    if re.fullmatch(r"\d+,\d+[eE]\+?\d+", s) or re.fullmatch(r"\d\.\d{1,9}[eE]\+?\d+", s):
        return "", "", "CNPJ em notação científica (dígitos perdidos) — redigite a célula como texto"
    s = _numero_excel(s)
    c = re.sub(r"[^0-9A-Z]", "", s.upper())
    if c.isdigit() and len(c) < 14:
        ok14 = cnpj_valido(c.zfill(14))
        ok11 = len(c) <= 11 and cpf_valido(c.zfill(11))
        if ok14 and ok11 and len(c) < 11:
            return c.zfill(14), "", "ambíguo: vale como CNPJ e como CPF com zeros à esquerda — informe com 14 ou 11 dígitos"
        if ok14 and not (len(c) == 11 and cpf_valido(c)):
            return c.zfill(14), "zeros à esquerda restaurados", ""
        if ok11 and len(c) != 11:
            return c.zfill(11), "zeros à esquerda restaurados", ""
    return c, "", ""


def tipo_inscricao(c: str):
    if len(c) == 14:
        return 1
    if len(c) == 11 and c.isdigit():
        return 2
    return None


def formatar_inscricao(c: str) -> str:
    if len(c) == 14:
        return f"{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}"
    if len(c) == 11:
        return f"{c[:3]}.{c[3:6]}.{c[6:9]}-{c[9:]}"
    return c


def para_data(v):
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = txt(v)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    if re.fullmatch(r"\d+(\.0)?", s) and 20000 < float(s) < 80000:      # número de série do Excel
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=int(float(s)))).date()
    return None


def para_int(v):
    s = digitos(_numero_excel(txt(v)))
    return int(s) if s else None


def zfill_se(v, n):
    d = digitos(_numero_excel(txt(v)))
    return d.zfill(n) if d and len(d) < n else d


# =====================================================================
# Referência (sem banco)
# =====================================================================
def _ler_csv(pasta, nome):
    with open(os.path.join(pasta, nome), encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f, delimiter=";"))


class Referencia:
    """Tabelas da Domínio exportadas por gerar_referencia.py."""

    def __init__(self, pasta):
        self.pasta = pasta
        linhas_cnae = _ler_csv(pasta, "cnae.csv")
        self.cnae = {r["cnae"]: r["descricao"] for r in linhas_cnae}
        self.cnae_prev = {r["cnae"] for r in linhas_cnae if r["cnae_esocial_antigo"] == "S"}
        self.cnae_fpas = {r["cnae"]: r for r in _ler_csv(pasta, "cnae_fpas.csv")}
        self.fpas = {int(r["fpas"]): r for r in _ler_csv(pasta, "fpas.csv")}
        self.terceiros = [{**r, "i_terceiros": int(r["i_terceiros"]), "fpas": int(r["fpas"]),
                           "codigo": int(r["codigo"]), "percentual": Decimal(r["percentual"])}
                          for r in _ler_csv(pasta, "terceiros.csv")]
        self.rat = {}
        for r in _ler_csv(pasta, "rat_cnae.csv"):
            self.rat.setdefault(r["cnae"], []).append((dt.date.fromisoformat(r["vigencia"]), Decimal(r["aliquota"])))
        self.municipios, self.por_srf, self.por_nome = {}, {}, {}
        for r in _ler_csv(pasta, "municipios.csv"):
            ibge = digitos(r["codigo_ibge"])
            if not ibge:
                continue
            self.municipios.setdefault(ibge, (int(r["codigo_municipio"]), r["nome"].strip(), r["uf"]))
            if digitos(r["codigo_srf"]):
                self.por_srf.setdefault(digitos(r["codigo_srf"]).lstrip("0"), ibge)
            self.por_nome.setdefault((norm(r["nome"]), r["uf"]), ibge)
        self.meta = {}
        for r in _ler_csv(pasta, "colunas.csv"):
            self.meta.setdefault(r["tabela"].upper(), []).append({
                "nome": r["coluna"], "up": r["coluna"].upper(), "tipo": r["tipo"], "largura": int(r["largura"] or 0),
                "escala": int(r["escala"] or 0), "nulo": r["nulo"] == "Y", "pk": r["pk"] == "Y",
                "padrao": r["padrao"] if r["padrao"] != "" else None})
        self.fks = {}
        for r in _ler_csv(pasta, "fks.csv"):
            self.fks.setdefault(r["tabela"].upper(), {})[r["coluna"].upper()] = r["tabela_mae"]
        try:
            with open(os.path.join(pasta, "versao.json"), encoding="utf-8") as f:
                self.versao = json.load(f)
        except (OSError, ValueError):
            self.versao = {}

    def municipio(self, ibge="", srf="", nome="", uf=""):
        """-> (código IBGE, (codigo_municipio, nome, uf)) pelo IBGE, código da Receita ou nome + UF."""
        for chave in (digitos(ibge), self.por_srf.get(digitos(srf).lstrip("0")),
                      self.por_nome.get((norm(nome), txt(uf).upper()))):
            if chave and chave in self.municipios:
                return chave, self.municipios[chave]
        return "", None

    def fpas_do_cnae(self, cnae):
        return self.cnae_fpas.get(zfill_se(cnae, 7), {})

    def terceiros_do_fpas(self, fpas):
        return sorted((t for t in self.terceiros if t["fpas"] == fpas), key=lambda t: t["codigo"])

    def terceiro_padrao(self, fpas, regra):
        opcoes = self.terceiros_do_fpas(fpas)
        if not opcoes:
            return None
        if regra == "zero":
            return next((t for t in opcoes if t["codigo"] == 0), opcoes[0])
        sem = [t for t in opcoes if norm(t["descricao"]) == "SEM CONVENIO"]
        return min(sem, key=lambda t: t["codigo"]) if sem else max(opcoes, key=lambda t: t["percentual"])

    def rat_do_cnae(self, cnae, data):
        vig = [a for v, a in self.rat.get(zfill_se(cnae, 7), []) if v <= data]
        return vig[-1] if vig else None


def ler_leiaute_txt(arquivos: dict):
    """Arquivos de leiaute exportados da Domínio ({nome: bytes}); uma linha por campo, separado por tab."""
    out = {}
    for _, dados in arquivos.items():
        conteudo = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                conteudo = dados.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        for linha in (conteudo or "").splitlines():
            p = linha.split("\t")
            if len(p) < 12 or not p[2].strip().isdigit():
                continue
            tab = p[1].strip().upper()
            if tab in TABELAS:
                out.setdefault(tab, {"linhas": [], "arquivo": ""})["linhas"].append(p[:12])
    for v in out.values():
        v["linhas"].sort(key=lambda r: int(str(r[2]).strip()))
    return out


def ler_leiaute_pasta(pasta):
    if not os.path.isdir(pasta):
        return {}
    arquivos = {}
    for n in os.listdir(pasta):
        if n.lower().endswith(".txt"):
            with open(os.path.join(pasta, n), "rb") as f:
                arquivos[n] = f.read()
    return ler_leiaute_txt(arquivos)


# =====================================================================
# Consultas externas: limite de vazão global, disjuntor e cache
# =====================================================================
class Provedor:
    """Limite de vazão compartilhado por todas as sessões do servidor (mesmo IP de saída)."""

    def __init__(self, nome, intervalo):
        self.nome, self.intervalo = nome, intervalo
        self.trava = threading.Lock()
        self.proximo, self.falhas, self.bloqueado_ate = 0.0, 0, 0.0

    def disponivel(self):
        return time.time() >= self.bloqueado_ate

    def aguardar(self):
        with self.trava:
            espera = self.proximo - time.time()
            if espera > 0:
                time.sleep(espera)
            self.proximo = time.time() + self.intervalo

    def sucesso(self):
        self.falhas = 0

    def falha(self, pausa=0):
        self.falhas += 1
        if pausa:
            self.proximo = max(self.proximo, time.time() + min(pausa, 120))
        if self.falhas >= 3:
            self.bloqueado_ate, self.falhas = time.time() + 300, 0


PROVEDORES = {"BrasilAPI": Provedor("BrasilAPI", 0.5), "CNPJá": Provedor("CNPJá", 12.5),
              "ReceitaWS": Provedor("ReceitaWS", 20.5), "ViaCEP": Provedor("ViaCEP", 0.2)}
_CACHE, _TRAVA_CACHE, TTL_CACHE = {}, threading.Lock(), 30 * 86400


class NaoEncontrado(Exception):
    pass


def _get_json(prov: Provedor, url, timeout=15):
    prov.aguardar()
    req = urllib.request.Request(url, headers={"User-Agent": "importador-tomadores-dominio/2.0",
                                               "Accept": "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            dados = json.loads(r.read().decode("utf-8"))
        prov.sucesso()
        return dados
    except urllib.error.HTTPError as e:
        if e.code in (400, 404):
            prov.sucesso()
            raise NaoEncontrado() from e
        espera = e.headers.get("Retry-After") if e.code == 429 else None
        prov.falha(int(espera) if espera and espera.isdigit() else (60 if e.code == 429 else 0))
        raise
    except (urllib.error.URLError, TimeoutError, ValueError):
        prov.falha()
        raise


def _cache(chave, funcao):
    with _TRAVA_CACHE:
        hit = _CACHE.get(chave)
        if hit and time.time() - hit[0] < TTL_CACHE:
            return hit[1]
    r = funcao()
    if not r.get("transitorio"):
        with _TRAVA_CACHE:
            _CACHE[chave] = (time.time(), r)
    return r


def _separar_tipo(logradouro):
    tipo, _, resto = txt(logradouro).partition(" ")
    return (norm(tipo), resto.strip()) if resto and norm(tipo) in TIPOS_LOGRADOURO else ("", txt(logradouro))


def _telefone(v):
    d = digitos(v)
    return (d[:2], d[2:]) if len(d) >= 10 else ("", d)


def _de_brasilapi(j):
    ddd, fone = _telefone(j.get("ddd_telefone_1"))
    regimes = sorted(j.get("regime_tributario") or [], key=lambda r: r.get("ano") or 0)
    return {"razao": txt(j.get("razao_social")), "fantasia": txt(j.get("nome_fantasia")),
            "cep": digitos(j.get("cep")), "tipo_logradouro": norm(j.get("descricao_tipo_de_logradouro")),
            "endereco": txt(j.get("logradouro")), "numero": txt(j.get("numero")),
            "complemento": txt(j.get("complemento")), "bairro": txt(j.get("bairro")),
            "cidade": txt(j.get("municipio")), "uf": txt(j.get("uf")), "ibge": digitos(j.get("codigo_municipio_ibge")),
            "srf": digitos(j.get("codigo_municipio")),
            "cnae": zfill_se(j.get("cnae_fiscal"), 7), "cnae_desc": txt(j.get("cnae_fiscal_descricao")),
            "ddd": ddd, "telefone": fone, "situacao": norm(j.get("descricao_situacao_cadastral")),
            "simples": j.get("opcao_pelo_simples"), "mei": j.get("opcao_pelo_mei"),
            "regime_receita": (f"{regimes[-1].get('forma_de_tributacao')} ({regimes[-1].get('ano')})"
                               if regimes else ""),
            "natureza": digitos(j.get("codigo_natureza_juridica")), "natureza_desc": txt(j.get("natureza_juridica")),
            "fonte": "BrasilAPI"}


def _de_cnpja(j):
    emp, end = j.get("company") or {}, j.get("address") or {}
    tipo, rua = _separar_tipo(end.get("street"))
    fone = (j.get("phones") or [{}])[0]
    ativ = j.get("mainActivity") or {}
    return {"razao": txt(emp.get("name")), "fantasia": txt(j.get("alias")), "cep": digitos(end.get("zip")),
            "tipo_logradouro": tipo, "endereco": rua, "numero": txt(end.get("number")),
            "complemento": txt(end.get("details")), "bairro": txt(end.get("district")), "cidade": txt(end.get("city")),
            "uf": txt(end.get("state")), "ibge": digitos(end.get("municipality")), "srf": "",
            "cnae": zfill_se(ativ.get("id"), 7), "cnae_desc": txt(ativ.get("text")),
            "ddd": digitos(fone.get("area")), "telefone": digitos(fone.get("number")),
            "situacao": norm((j.get("status") or {}).get("text")),
            "simples": (emp.get("simples") or {}).get("optant"), "mei": (emp.get("simei") or {}).get("optant"),
            "regime_receita": "", "natureza": digitos((emp.get("nature") or {}).get("id")),
            "natureza_desc": txt((emp.get("nature") or {}).get("text")), "fonte": "CNPJá"}


def _de_receitaws(j):
    tipo, rua = _separar_tipo(j.get("logradouro"))
    ddd, fone = _telefone(txt(j.get("telefone")).split("/")[0])
    ativ = (j.get("atividade_principal") or [{}])[0]
    nat = txt(j.get("natureza_juridica"))
    return {"razao": txt(j.get("nome")), "fantasia": txt(j.get("fantasia")), "cep": digitos(j.get("cep")),
            "tipo_logradouro": tipo, "endereco": rua, "numero": txt(j.get("numero")),
            "complemento": txt(j.get("complemento")), "bairro": txt(j.get("bairro")),
            "cidade": txt(j.get("municipio")), "uf": txt(j.get("uf")), "ibge": "", "srf": "",
            "cnae": digitos(ativ.get("code")), "cnae_desc": txt(ativ.get("text")), "ddd": ddd, "telefone": fone,
            "situacao": norm(j.get("situacao")), "simples": (j.get("simples") or {}).get("optante"),
            "mei": (j.get("simei") or {}).get("optante"), "regime_receita": "",
            "natureza": digitos(nat.split(" - ")[0]), "natureza_desc": nat.partition(" - ")[2], "fonte": "ReceitaWS"}


def consultar_cnpj(cnpj: str) -> dict:
    """BrasilAPI -> CNPJá -> ReceitaWS. Erro transitório (limite, fora do ar) não vai para o cache."""
    cnpj = re.sub(r"[^0-9A-Z]", "", txt(cnpj).upper())

    def buscar():
        tentativas = [("BrasilAPI", f"https://brasilapi.com.br/api/cnpj/v1/{cnpj}", _de_brasilapi),
                      ("CNPJá", f"https://open.cnpja.com/office/{cnpj}", _de_cnpja),
                      ("ReceitaWS", f"https://receitaws.com.br/v1/cnpj/{cnpj}", _de_receitaws)]
        falhas = []
        for nome, url, conv in tentativas:
            prov = PROVEDORES[nome]
            if not prov.disponivel():
                falhas.append(f"{nome} em pausa")
                continue
            try:
                j = _get_json(prov, url)
            except NaoEncontrado:
                return {"erro": "CNPJ não encontrado na Receita"}
            except (urllib.error.URLError, TimeoutError, ValueError) as e:
                falhas.append(f"{nome}: {getattr(e, 'code', '') or e.__class__.__name__}")
                continue
            if nome == "ReceitaWS" and txt(j.get("status")).upper() == "ERROR":
                return {"erro": txt(j.get("message")) or "CNPJ não encontrado"}
            r = conv(j)
            if not r["ibge"] and not r["srf"] and r["cep"]:
                r["ibge"] = consultar_cep(r["cep"]).get("ibge", "")
            return r
        return {"erro": "consulta indisponível agora (" + "; ".join(falhas) + ") — tente de novo",
                "transitorio": True}

    return _cache(("cnpj", cnpj), buscar)


def consultar_cep(cep: str) -> dict:
    cep = zfill_se(cep, 8)
    if len(cep) != 8:
        return {"erro": "CEP inválido"}

    def buscar():
        try:
            j = _get_json(PROVEDORES["ViaCEP"], f"https://viacep.com.br/ws/{cep}/json/")
        except NaoEncontrado:
            return {"erro": "CEP não encontrado"}
        except (urllib.error.URLError, TimeoutError, ValueError) as e:
            return {"erro": f"falha na consulta do CEP ({e.__class__.__name__})", "transitorio": True}
        if j.get("erro"):
            return {"erro": "CEP não encontrado"}
        tipo, rua = _separar_tipo(j.get("logradouro"))
        return {"cep": cep, "tipo_logradouro": tipo, "endereco": rua, "bairro": txt(j.get("bairro")),
                "cidade": txt(j.get("localidade")), "uf": txt(j.get("uf")), "ibge": digitos(j.get("ibge"))}

    return _cache(("cep", cep), buscar)


# =====================================================================
# Perfil tributário da prestadora (herdado por todos os tomadores)
# =====================================================================
def sugerir_regime(rec: dict):
    """-> (chave do regime, motivo)."""
    if rec.get("mei"):
        return "MEI", "optante pelo SIMEI na Receita"
    if rec.get("simples"):
        if txt(rec.get("cnae")).startswith(ANEXO_IV_PREFIXOS):
            return "SN4", "optante pelo Simples; CNAE principal de atividade do Anexo IV"
        return "SN", "optante pelo Simples — confirme o anexo (cessão de mão de obra costuma ser Anexo IV)"
    if rec.get("simples") is False:
        return "LRLP", "não optante pelo Simples" + (f" — {rec['regime_receita']}" if rec.get("regime_receita") else "")
    return "LRLP", "regime não informado pela consulta — confirme"


def sugerir_perfil(rec: dict, ref: Referencia, data: dt.date) -> dict:
    """Sugestões a partir do CNAE principal e do regime da prestadora; tudo pode ser ajustado na tela."""
    regime, motivo_regime = sugerir_regime(rec)
    lig = ref.fpas_do_cnae(rec.get("cnae"))
    fpas = int(lig["fpas"]) if txt(lig.get("fpas")) else None
    alternativas = [int(x) for x in txt(lig.get("alternativas")).split("|") if x.strip().isdigit()]
    alertas = []
    if rec.get("situacao") and rec["situacao"] != "ATIVA":
        alertas.append(f"Situação cadastral da prestadora na Receita: {rec['situacao']}.")
    if not lig:
        alertas.append("CNAE principal não está na tabela de CNAE da Domínio — escolha o FPAS.")
    elif lig["ligacao"] == "ambígua":
        alertas.append(f"O CNAE {lig['cnae']} aparece com mais de um FPAS ({', '.join(map(str, alternativas))}) — "
                       "escolha pela atividade preponderante.")
    elif lig["ligacao"] == "sem FPAS":
        alertas.append(f"O CNAE {lig['cnae']} não tem FPAS na relação da Domínio — escolha o FPAS.")
    elif lig["ligacao"] != "exata":
        alertas.append(f"FPAS {fpas} ligado ao CNAE {lig['ligacao']} — confira.")
    nat = txt(rec.get("natureza"))
    if nat.startswith("1"):
        alertas.append("Natureza jurídica de administração pública: normalmente FPAS 582.")
    elif nat.startswith("3"):
        alertas.append("Entidade sem fins lucrativos: se for beneficente com CEBAS, o FPAS é 639 (isenção).")
    rat = ref.rat_do_cnae(rec.get("cnae"), data)
    if rat is None:
        alertas.append("Sem alíquota RAT para o CNAE principal na tabela da Domínio — informe o RAT.")
    return {"regime": regime, "motivo_regime": motivo_regime, "fpas": fpas, "alternativas": alternativas,
            "ligacao": lig.get("ligacao", ""), "i_fpas": int(lig["i_fpas"]) if txt(lig.get("i_fpas")) else None,
            "atividade": lig.get("atividade_fofpas", ""), "rat": rat, "alertas": alertas}


def encargos_do_perfil(p: dict) -> dict:
    """Perfil escolhido na tela -> colunas da vigência do tomador (iguais para todos)."""
    t = p["terceiro"]
    return {"CODIGO_FPAS": p["fpas"], "I_FPAS": p.get("i_fpas"),
            "CODIGO_TERCEIRO": t["codigo"], "I_TERCEIROS": t["i_terceiros"],
            "PERC_TERCEIRO": t["percentual"], "PERCENTUAL_TERCEIRO_BRUTO": t["percentual"],
            "PERC_INSS_EMPRESA": Decimal(str(p["inss"])), "PERC_AUTONOMOS": Decimal(str(p["prolab"])),
            "PERC_ACID_TRABALHO": Decimal(str(p["rat"])), "TIPO_INFORMACAO_ALIQUOTA_ACIDENTE_TRABALHO": RAT_INFORMADO,
            "CODIGO_SAT": 0, "FILANTROPIA": Decimal(0), "CALCULA_INSS_EMPRESA_LEI_12546": 0, "SEQUENCIA_GPS": 1,
            "NUMERO_FGTS": 0, "INEXIGIBILIDADE_RAT": 0, "CODIGO_GFIP": p["sefip"], "CODIGO_GPS": p["gps"]}


# =====================================================================
# Tomadores (estrutura única para digitação e planilha)
# =====================================================================
CAMPOS = [  # chave, rótulo da planilha, ajuda
    ("inscricao", "CNPJ/CPF do tomador", "Obrigatório. Com ou sem pontuação (CNPJ alfanumérico aceito)."),
    ("nome", "Nome do serviço", "Até 40 caracteres. Vazio: razão social da Receita."),
    ("data_inicio", "Data de início", "DD/MM/AAAA. Vazio: data informada na tela."),
    ("cep", "CEP", "8 dígitos. Com CEP, o app completa endereço e município."),
    ("tipo_logradouro", "Tipo de logradouro", "Ex.: RUA, AVENIDA."),
    ("endereco", "Endereço", "Até 31 caracteres."),
    ("numero", "Número", "Somente números (S/N: deixe vazio)."),
    ("complemento", "Complemento", "Até 30 caracteres."),
    ("bairro", "Bairro", "Até 20 caracteres."),
    ("cidade", "Cidade", "Até 20 caracteres."),
    ("uf", "UF", "Sigla do estado."),
    ("ibge", "Código IBGE do município", "7 dígitos. Vazio: buscado pelo CEP ou pela cidade + UF."),
    ("cnae", "CNAE 2.3", "7 dígitos (ex.: 8112500). Vazio: CNAE principal na Receita."),
    ("ddd", "DDD", ""),
    ("telefone", "Telefone", ""),
    ("tipo_servico", "Tipo de serviço do tomador", "1, 2 ou 3. Vazio: padrão da tela."),
    ("lotacao", "Tipo de lotação eSocial", "Código gravado na Domínio. Vazio: padrão da tela."),
]
ROTULO = {k: r for k, r, _ in CAMPOS}
SINONIMOS = {
    "inscricao": {"CNPJ CPF DO TOMADOR", "CNPJ CPF", "CNPJ", "CPF", "INSCRICAO", "CNPJ DO TOMADOR", "CPF CNPJ"},
    "nome": {"NOME DO SERVICO", "NOME", "RAZAO SOCIAL", "TOMADOR", "CLIENTE"},
    "data_inicio": {"DATA DE INICIO", "DATA INICIO", "INICIO", "VIGENCIA"},
    "tipo_logradouro": {"TIPO DE LOGRADOURO", "TIPO LOGRADOURO", "TIPO ENDERECO"},
    "endereco": {"ENDERECO", "LOGRADOURO", "RUA"},
    "numero": {"NUMERO", "NO", "NRO", "N"},
    "ibge": {"CODIGO IBGE DO MUNICIPIO", "CODIGO IBGE", "IBGE", "COD IBGE"},
    "cnae": {"CNAE 2 3", "CNAE", "CNAE PRINCIPAL"},
    "tipo_servico": {"TIPO DE SERVICO DO TOMADOR", "TIPO SERVICO TOMADOR", "TIPO DE SERVICO"},
    "lotacao": {"TIPO DE LOTACAO ESOCIAL", "LOTACAO", "TIPO LOTACAO"},
    "cidade": {"CIDADE", "MUNICIPIO"},
}
EXTRAS = ("razao", "situacao", "fonte", "erro", "aviso_insc", "erro_insc", "srf")


def tomador_vazio(valor="", linha=None) -> dict:
    ins, aviso, erro = normalizar_inscricao(valor)
    t = {k: "" for k, _, _ in CAMPOS}
    t.update({e: "" for e in EXTRAS})
    t.update({"inscricao": ins, "aviso_insc": aviso, "erro_insc": erro, "linha": linha, "original": txt(valor),
              "consultado": False})
    return t


def completar(t: dict, consultar=True) -> dict:
    """Preenche só os campos vazios: Receita (CNPJ) e, se ainda faltar endereço ou município, ViaCEP."""
    t = dict(t)
    ins = t["inscricao"]
    t["consultado"] = True
    if consultar and not t["erro_insc"] and tipo_inscricao(ins) == 1 and cnpj_valido(ins):
        r = consultar_cnpj(ins)
        if r.get("erro"):
            t["erro"] = r["erro"]
            t["consultado"] = not r.get("transitorio")
        else:
            t.update({"razao": r["razao"], "situacao": r["situacao"], "fonte": r["fonte"], "srf": r["srf"],
                      "erro": ""})
            if not t["nome"]:
                t["nome"] = r["razao"]
            for k in ("cep", "tipo_logradouro", "endereco", "numero", "complemento", "bairro", "cidade", "uf",
                      "ibge", "cnae", "ddd", "telefone"):
                if not txt(t.get(k)) and r.get(k):
                    t[k] = r[k]
    if consultar and zfill_se(t["cep"], 8) and (not digitos(t["ibge"]) or not txt(t["endereco"])):
        r = consultar_cep(t["cep"])
        if not r.get("erro"):
            for k in ("tipo_logradouro", "endereco", "bairro", "cidade", "uf", "ibge"):
                if not txt(t.get(k)) and r.get(k):
                    t[k] = r[k]
            t["fonte"] = (t["fonte"] + " + ViaCEP").strip(" +")
    return t


def _grade(dados: bytes, nome: str) -> pd.DataFrame:
    if nome.lower().endswith((".csv", ".txt")):
        conteudo = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                conteudo = dados.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        try:
            delim = csv.Sniffer().sniff(conteudo[:5000], delimiters=";,\t|").delimiter
        except csv.Error:
            delim = ";"
        return pd.read_csv(io.StringIO(conteudo), sep=delim, dtype=str, header=None, keep_default_na=False)
    return pd.read_excel(io.BytesIO(dados), dtype=str, header=None, keep_default_na=False)


def ler_planilha(dados: bytes, nome: str):
    """Planilha do cliente -> (tomadores, info). Lê tudo como texto, acha o cabeçalho, ignora linhas em branco
    e guarda a linha real do Excel de cada registro."""
    bruto = _grade(dados, nome)
    cab, mapa = None, {}
    for i in range(min(25, len(bruto))):
        tentativa = {}
        for j, v in enumerate(bruto.iloc[i].tolist()):
            n = norm(v)
            for k, r, _ in CAMPOS:
                if k not in tentativa.values() and n and (n == norm(r) or n in SINONIMOS.get(k, set())):
                    tentativa[j] = k
                    break
        if "inscricao" in tentativa.values():
            cab, mapa = i, tentativa
            break
    if cab is None:
        raise ValueError("Não encontrei a coluna 'CNPJ/CPF do tomador' nas primeiras 25 linhas.")
    nao_lidas = [txt(v) for j, v in enumerate(bruto.iloc[cab].tolist()) if j not in mapa and txt(v)]
    tomadores, em_branco = [], 0
    j_ins = next(j for j, k in mapa.items() if k == "inscricao")
    for i in range(cab + 1, len(bruto)):
        linha = bruto.iloc[i].tolist()
        if not any(txt(v) for v in linha):
            em_branco += 1
            continue
        if norm(next(txt(v) for v in linha if txt(v))).startswith("TOTAL"):
            continue
        if not txt(linha[j_ins]):
            em_branco += 1
            continue
        t = tomador_vazio(linha[j_ins], linha=i + 1)
        for j, k in mapa.items():
            if k != "inscricao":
                t[k] = txt(linha[j])
        t["cep"] = zfill_se(t["cep"], 8)
        t["cnae"] = zfill_se(t["cnae"], 7)
        t["ibge"] = digitos(_numero_excel(t["ibge"]))
        t["fonte"] = "Planilha"
        tomadores.append(t)
    return tomadores, {"cabecalho": cab + 1, "em_branco": em_branco, "nao_lidas": nao_lidas,
                       "colunas": sorted(set(mapa.values()))}


def planilha_resultado(dados: bytes, nome: str, linha_cab: int, resultado: dict) -> bytes:
    """Devolve a planilha do cliente com as colunas Status e Motivo nas mesmas linhas."""
    bruto = _grade(dados, nome)
    status = [("Status" if i + 1 == linha_cab else resultado.get(i + 1, ("", ""))[0]) for i in range(len(bruto))]
    motivo = [("Motivo" if i + 1 == linha_cab else resultado.get(i + 1, ("", ""))[1]) for i in range(len(bruto))]
    bruto["Status"], bruto["Motivo"] = status, motivo
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        bruto.to_excel(w, sheet_name="resultado", index=False, header=False)
        for c in w.sheets["resultado"][linha_cab]:
            c.font = Font(bold=True)
    return buf.getvalue()


def modelo_planilha() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "tomadores"
    ws.append([r for _, r, _ in CAMPOS])
    for i, (k, _, _) in enumerate(CAMPOS, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78" if k == "inscricao" else "4472C4")
        ws.column_dimensions[c.column_letter].width = max(14, len(c.value) + 4)
        for r in range(2, 1002):
            ws.cell(row=r, column=i).number_format = "@"
    ws.freeze_panes = "A2"
    ins = wb.create_sheet("instrucoes")
    ins.append(["Coluna", "Como preencher"])
    for _, r, h in CAMPOS:
        ins.append([r, h])
    ins.append([])
    ins.append(["Obs.", "Só o CNPJ/CPF é obrigatório. Com CNPJ, o app completa pela Receita os campos vazios. "
                        "FPAS, terceiros, RAT, SEFIP e GPS NÃO vão na planilha: vêm da empresa prestadora."])
    ins.append(["Versão do modelo", "2"])
    for c in ins[1]:
        c.font = Font(bold=True)
    ins.column_dimensions["A"].width, ins.column_dimensions["B"].width = 32, 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# =====================================================================
# Validação de cada tomador
# =====================================================================
def validar(t: dict, ctx: dict) -> dict:
    """-> {'status', 'obs', 'dados'}; ctx traz a referência, os padrões da tela e os CNPJs já cadastrados."""
    ref = ctx["ref"]
    erros, avisos = [], []
    ins = t.get("inscricao", "")
    tipo = tipo_inscricao(ins)
    if t.get("erro_insc"):
        erros.append(t["erro_insc"])
    elif tipo == 1 and not cnpj_valido(ins):
        erros.append("CNPJ inválido (dígito verificador não confere)")
    elif tipo == 2 and not cpf_valido(ins):
        erros.append("CPF inválido (dígito verificador não confere)")
    elif tipo is None:
        erros.append("inscrição deve ter 14 (CNPJ) ou 11 (CPF) caracteres")
    if t.get("aviso_insc"):
        avisos.append(t["aviso_insc"])
    if txt(t.get("erro")):
        (erros if "não encontrado" in t["erro"].lower() else avisos).append(f"Receita: {t['erro']}")
    sit = txt(t.get("situacao"))
    if sit and sit != "ATIVA":
        avisos.append(f"situação na Receita: {sit}")
    if ins and ins == ctx.get("cnpj_prestadora"):
        erros.append("é o CNPJ da própria prestadora")
    if ins in ctx["cadastrados"]:
        erros.append("já cadastrado como serviço na empresa")

    nome = txt(t.get("nome"))
    if not nome:
        erros.append("nome do serviço vazio")
    data = para_data(t.get("data_inicio")) or ctx["data_padrao"]
    if txt(t.get("data_inicio")) and not para_data(t.get("data_inicio")):
        erros.append(f"data de início inválida: {t['data_inicio']}")

    cnae = zfill_se(t.get("cnae"), 7) if ctx["cnae_cadastro"] == "tomador" else ctx["cnae_prestadora"]
    if not cnae:
        erros.append("CNAE vazio")
    elif cnae not in ref.cnae:
        erros.append(f"CNAE {cnae} não existe na tabela de CNAE 2.3 da Domínio")

    ibge, mun = ref.municipio(t.get("ibge"), t.get("srf"), t.get("cidade"), t.get("uf"))
    if not mun:
        erros.append("município não encontrado — informe o CEP, o código IBGE ou cidade + UF")

    cep = zfill_se(t.get("cep"), 8)
    if cep and len(cep) != 8:
        avisos.append(f"CEP com {len(cep)} dígitos")
    tipo_lg = norm(t.get("tipo_logradouro"))
    endereco = txt(t.get("endereco"))
    tipo_end = ctx["de_para_logradouro"].get(tipo_lg)
    if tipo_lg and tipo_end is None and not norm(endereco).startswith(tipo_lg):
        endereco = f"{tipo_lg} {endereco}".strip()
    if not endereco:
        avisos.append("endereço vazio")
    numero_txt = _numero_excel(txt(t.get("numero")))
    numero = int(numero_txt) if re.fullmatch(r"\d+", numero_txt or "") else None
    complemento = txt(t.get("complemento"))
    if numero_txt and numero is None:
        complemento = f"{numero_txt} {complemento}".strip()

    tipo_serv = para_int(t.get("tipo_servico")) or ctx["tipo_servico"]
    if tipo_serv not in TIPOS_SERVICO_TOMADOR:
        erros.append(f"tipo de serviço do tomador {tipo_serv} inválido (1, 2 ou 3)")
    lot = para_int(t.get("lotacao")) or (ctx["lotacao_cpf"] if tipo == 2 else ctx["lotacao_cnpj"])

    dados = {"inscricao": ins, "tipo_insc": tipo, "nome": nome, "data": data, "cep": cep or None,
             "tipo_endereco": tipo_end, "endereco": endereco or None, "numero": numero,
             "complemento": complemento or None, "bairro": txt(t.get("bairro")) or None,
             "cidade": txt(t.get("cidade")) or (mun[1] if mun else None),
             "uf": (txt(t.get("uf")).upper() or (mun[2] if mun else "")) or None,
             "codigo_municipio": mun[0] if mun else None, "ibge": ibge, "cnae": cnae or None,
             "ddd": digitos(t.get("ddd")) or None, "telefone": digitos(t.get("telefone")) or None,
             "tipo_servico": tipo_serv, "lotacao": lot}
    for campo, lim in ctx["limites"].items():
        v = dados.get(campo)
        if isinstance(v, str) and lim and len(v) > lim:
            avisos.append(f"{campo} cortado em {lim} caracteres")
            dados[campo] = v[:lim].rstrip()
    status = "❌ Erro" if erros else ("⚠️ Revisar" if avisos else "✅ Pronto")
    return {"status": status, "obs": " | ".join(erros + avisos), "dados": dados}


# =====================================================================
# Montagem das linhas das tabelas
# =====================================================================
def _vazio_padrao(col):
    return "" if col["tipo"] in TEXTOS else (None if col["tipo"] in DATAS else 0)


# Campos que o app só preenche em alguns casos de configuração (de-para de logradouro; vigência anterior a 12/2018).
# Vazios no lote inteiro, saem do leiaute; os demais campos atribuídos ficam sempre no leiaute (leiaute estável).
OPCIONAIS = {"TIPO_ENDERECO", "I_CNAE_ESOCIAL"}


def montar_linhas(itens, encargos: dict, cfg: dict, meta: dict):
    """itens: [(i_servicos, dados validados)]. Todos herdam os mesmos encargos da prestadora.
    Retorna (linhas por tabela, colunas que o app atribui em cada tabela)."""
    serv, vig = [], []
    atrib = {T_SERV: set(), T_VIG: set()}
    for cod, d in itens:
        v = {c["up"]: None for c in meta[T_VIG]}
        dados_v = {k: val for k, val in encargos.items() if k in v}
        dados_v.update({
            "CODI_EMP": cfg["empresa"], "I_SERVICOS": cod, "VIGENCIA": d["data"], "DESCRICAO": cfg["descricao"],
            "TIPO": TIPO_TOMADOR, "ORIGEM_REG": ORIGEM_IMP_TABELAS, "ORIGEM_ALTERACAO": 1,
            "COMPETENCIA_FIM_VIGENCIA": FIM_ABERTO, "I_FILIAIS": cfg["filial"],
            "TIPO_LOTACAO_ESOCIAL": d["lotacao"], "TIPO_SERVICO_TOMADOR": d["tipo_servico"],
            "ENVIAR_ESOCIAL_S_1020": 1,
            "CGC": d["inscricao"], "TIPO_INSC": d["tipo_insc"], "ENDERECO": d["endereco"], "NUMERO": d["numero"],
            "BAIRRO": d["bairro"], "CEP": d["cep"], "CIDADE": d["cidade"], "ESTADO": d["uf"],
            "CODIGO_MUNICIPIO": d["codigo_municipio"], "TIPO_ENDERECO": d["tipo_endereco"],
            "COMPLEMENTO": d["complemento"], "DDD": d["ddd"], "TELEFONE": d["telefone"],
            "I_CNAE20": d["cnae"], "CODIGO_ATIVIDADE": 0,
            "I_CNAE_ESOCIAL": d["cnae"] if (d["data"] < CORTE_CNAE_ESOCIAL and d["cnae"] in cfg["cnae_prev"]) else None,
        })
        v.update(dados_v)
        s = {c["up"]: None for c in meta[T_SERV]}
        dados_s = {"CODI_EMP": cfg["empresa"], "I_SERVICOS": cod, "NOME": d["nome"], "ORIGEM_REG": ORIGEM_IMP_TABELAS,
                   "DATA_INICIO": d["data"], "SITUACAO": 1, "ENVIAR_ESOCIAL": 1, "CODIGO_ESOCIAL": str(cod)}
        s.update(dados_s)
        atrib[T_VIG].update(k for k in dados_v if k in v)
        atrib[T_SERV].update(dados_s)
        for linha, tab in ((s, T_SERV), (v, T_VIG)):
            for c in meta[tab]:
                if linha.get(c["up"]) is None and not c["nulo"] and c["padrao"] is None:
                    linha[c["up"]] = _vazio_padrao(c)
                    atrib[tab].add(c["up"])
        serv.append(s)
        vig.append(v)
    return {T_SERV: serv, T_VIG: vig}, atrib


# =====================================================================
# Arquivos da importação
# =====================================================================
def formatar(v, col) -> str:
    """Valor -> texto do TXT da importação (decimais embutidos, como no Acumuladores.py)."""
    if v is None:
        return ""
    t = col["tipo"]
    if t in DATAS:
        return (v.date() if isinstance(v, dt.datetime) else v).strftime("%d/%m/%Y")
    if t in DECIMAIS:
        d = Decimal(str(v))
        if col["escala"] > 0:
            d = d * (Decimal(10) ** col["escala"])
        return str(int(d.quantize(Decimal(1), ROUND_HALF_UP)))
    if t in INTEIROS:
        return str(int(v))
    return str(v).replace("\t", " ").replace("\r", " ").replace("\n", " ")


def _campos_leiaute(lei_tab):
    return [(r, str(r[3]).strip().upper(), (str(r[9]).strip() or "I").upper(), str(r[10]).strip())
            for r in lei_tab["linhas"]]


def gerar_pacote(linhas: dict, atrib: dict, meta: dict, leiaute: dict, fks: dict, ctx: dict, destino,
                 perfil_txt="", empresa=None):
    """Monta o .zip (dados + leiaute ajustado + estrutura + leia-me), a lista de conferências e as informações
    para as instruções de importação. O leiaute ajustado é estável: só depende do que o app preenche."""
    checks, arquivos, estrutura, removidos = [], {}, [], {}
    info = {"destino": destino, "arquivos": {}, "campos": {}, "empresa": empresa,
            "codigos": [r["I_SERVICOS"] for r in linhas[T_SERV]]}

    def add(nome, nivel, problemas):
        ok = not problemas
        checks.append({"Verificação": nome, "Nível": nivel,
                       "Resultado": "✅" if ok else ("❌" if nivel == "Erro" else "⚠️"),
                       "Detalhe": "ok" if ok else "; ".join(problemas)[:900]})

    opcionais = OPCIONAIS - set(ctx.get("manter", ()))
    for tab in TABELAS:
        cols = {c["up"]: c for c in meta[tab]}
        rows = linhas[tab]
        campos = _campos_leiaute(leiaute[tab])
        nomes_lei = {cp for _, cp, _, _ in campos}
        add(f"{tab}: leiaute encontrado", "Erro", [] if campos else ["leiaute sem campos para esta tabela"])
        pk_fora = [c["nome"] for c in meta[tab] if c["pk"] and c["up"] not in nomes_lei]
        add(f"{tab}: chave primária no leiaute", "Erro", [f"falta {', '.join(pk_fora)}"] if pk_fora else [])

        tirar, prob_nn, mistos_fk, mistos = set(), [], [], []
        for _, cp, origem, _ in campos:
            c = cols.get(cp)
            if not c or origem == "F":
                continue
            vals = [r.get(cp) for r in rows]
            vazios = sum(v is None for v in vals)
            if cp not in atrib[tab] or (vazios == len(vals) and (cp in opcionais or cp in fks.get(tab, {}))):
                if c["nulo"] or c["padrao"] is not None:
                    tirar.add(cp)
                else:
                    prob_nn.append(c["nome"])
            elif vazios:
                (mistos_fk if cp in fks.get(tab, {}) else mistos).append(f"{c['nome']} ({vazios} vazio(s))")
        removidos[tab] = tirar
        fora = [c["nome"] for c in meta[tab] if not c["nulo"] and c["padrao"] is None and c["up"] not in nomes_lei]
        add(f"{tab}: campos obrigatórios sem padrão estão no leiaute", "Erro",
            [f"fora do leiaute: {', '.join(fora)}"] if fora else [])
        add(f"{tab}: campos obrigatórios preenchidos", "Erro", [f"vazio: {', '.join(prob_nn)}"] if prob_nn else [])
        add(f"{tab}: chave estrangeira vazia só em parte das linhas (importador gravaria 0)", "Erro",
            [", ".join(mistos_fk) + " — gere em lotes separados"] if mistos_fk else [])
        add(f"{tab}: campos vazios só em parte das linhas (importador grava 0 ou '')", "Aviso",
            [", ".join(mistos)] if mistos else [])
        fixos = [f"{cp}={fixo or '(vazio)'}" for _, cp, origem, fixo in campos
                 if origem == "F" and cp in cols and any(r.get(cp) is not None for r in rows)]
        add(f"{tab}: campos com valor fixo no leiaute (o valor do app não vai)", "Aviso",
            [", ".join(fixos)] if fixos else [])
        ignorados = [cp for _, cp, origem, _ in campos if cp not in cols and origem != "F"]
        add(f"{tab}: campos do leiaute que não são coluna da tabela (importador ignora)", "Aviso",
            [", ".join(ignorados)] if ignorados else [])

        larg, rng = [], []
        for r in rows:
            for cp, v in r.items():
                c = cols.get(cp)
                if v is None or not c:
                    continue
                if c["tipo"] in TEXTOS and c["largura"] and len(str(v)) > c["largura"]:
                    larg.append(f"{c['nome']}={str(v)[:20]}…")
                if c["tipo"] in LIMITES_INT:
                    lo, hi = LIMITES_INT[c["tipo"]]
                    if not lo <= int(v) <= hi:
                        rng.append(f"{c['nome']}={v}")
                if c["tipo"] in DECIMAIS and c["largura"]:
                    if len(str(abs(int(Decimal(str(v)))))) > c["largura"] - c["escala"]:
                        rng.append(f"{c['nome']}={v}")
        add(f"{tab}: textos dentro do tamanho da coluna", "Erro", larg[:20])
        add(f"{tab}: números dentro do limite da coluna", "Erro", rng[:20])
        pk = [c["up"] for c in meta[tab] if c["pk"]]
        chaves = [tuple(r.get(k) for k in pk) for r in rows]
        dup = len(chaves) - len(set(chaves))
        add(f"{tab}: chave única no arquivo", "Erro", [f"{dup} chave(s) repetida(s)"] if dup else [])

        mantidos = [(r, cp, origem) for r, cp, origem, _ in campos if cp not in tirar]
        lidos = [(r, cp) for r, cp, origem in mantidos if origem != "F"]
        conteudo = ["\t".join(formatar(lin.get(cp), cols[cp]) if cp in cols else "" for _, cp in lidos)
                    for lin in rows]
        nome_arq = leiaute[tab].get("arquivo") or f"{tab}.txt"
        if not nome_arq.lower().endswith(".txt"):
            nome_arq += ".txt"
        arquivos[f"dados/{nome_arq}"] = ("\r\n".join(conteudo) + "\r\n").encode("cp1252", errors="replace")
        info["arquivos"][tab], info["campos"][tab] = nome_arq, len(mantidos)
        lei_lin = []
        for seq, (r, cp, origem) in enumerate(mantidos, 1):
            vals = list(r)
            vals[0], vals[2] = destino, seq
            lei_lin.append("\t".join("" if v is None else str(v) for v in vals))
        arquivos[f"leiaute_ajustado/{tab}.txt"] = ("\r\n".join(lei_lin) + "\r\n").encode("cp1252", errors="replace")
        for p, (r, cp) in enumerate(lidos, 1):
            c = cols.get(cp)
            estrutura.append([nome_arq, p, r[2], cp, (c or {}).get("tipo", ""), (c or {}).get("largura", ""),
                              (c or {}).get("escala", ""), "S" if c and c["pk"] else "N",
                              "S" if c and c["nulo"] else "N", "" if c else "ignorado pelo importador"])

    sem_serv = {(r["CODI_EMP"], r["I_SERVICOS"]) for r in linhas[T_VIG]} - \
               {(r["CODI_EMP"], r["I_SERVICOS"]) for r in linhas[T_SERV]}
    add("FOVIGENCIAS_SERVICO → FOSERVICOS: toda vigência tem o serviço no arquivo", "Erro",
        [f"{len(sem_serv)} vigência(s) sem serviço"] if sem_serv else [])
    m_ruins = [r["CGC"] for r in linhas[T_VIG] if r.get("CODIGO_MUNICIPIO") not in ctx["codigos_municipio"]]
    add("Município existe na Domínio", "Erro", [f"{len(m_ruins)} linha(s)"] if m_ruins else [])
    c_ruins = [r["CGC"] for r in linhas[T_VIG] if r.get("I_CNAE20") not in ctx["cnae20"]]
    add("CNAE existe na Domínio", "Erro", [f"{len(c_ruins)} linha(s)"] if c_ruins else [])
    t_ruins = [r["CGC"] for r in linhas[T_VIG] if r.get("I_TERCEIROS") not in ctx["terceiros_ids"]]
    add("Código de terceiros existe na Domínio para o FPAS", "Erro", [f"{len(t_ruins)} linha(s)"] if t_ruins else [])

    with io.StringIO() as s:
        w = csv.writer(s, delimiter=";")
        w.writerow(["arquivo", "posicao", "seq_leiaute", "coluna", "tipo", "tamanho", "decimais", "chave",
                    "aceita_nulo", "obs"])
        w.writerows(estrutura)
        arquivos["estrutura_arquivos.csv"] = s.getvalue().encode("utf-8-sig")
    tirados = {t: sorted(v) for t, v in removidos.items() if v}
    info["assinatura"] = hashlib.md5("|".join(
        f"{t}:{','.join(cp for _, cp, _, _ in _campos_leiaute(leiaute[t]) if cp not in removidos[t])}"
        for t in TABELAS).encode()).hexdigest()[:8].upper()
    arquivos["LEIA-ME.txt"] = leia_me(info, perfil_txt).encode("cp1252", errors="replace")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for nome, conteudo in arquivos.items():
            z.writestr(nome, conteudo)
    return buf.getvalue(), checks, tirados, info


def _faixa(cods):
    a, b = min(cods), max(cods)
    return f"código {a}" if a == b else f"códigos {a} a {b}"


def _sem_negrito(s):
    return s.replace("**", "")


def passos_importacao(info: dict):
    """Passo a passo da importação na Domínio, em linguagem simples (usado na tela e no LEIA-ME).
    Cada etapa: {'icone', 'titulo', 'quando', 'passos', 'dica'}. O **negrito** vira texto simples no LEIA-ME."""
    lei, arq, n = info["destino"], info["arquivos"], info["campos"]
    emp, cods = info.get("empresa"), info["codigos"]
    a, b = min(cods), max(cods)
    livre = f"o código {a} está livre" if a == b else f"os códigos {a} a {b} estão livres"
    return [
        {"icone": "🧩", "titulo": "Etapa 1 · Preparar o leiaute",
         "quando": f"Só na primeira vez. Se você já importou o leiaute com o código {info['assinatura']}, "
                   "pule direto para a etapa 2.",
         "passos": [
             f"Na Domínio, abra a importação por leiaute e escolha o leiaute **{lei}**.",
             "Apague os campos que ele já tiver nas tabelas FOSERVICOS e FOVIGENCIAS_SERVICO.",
             "Importe os dois arquivos da pasta **leiaute_ajustado** do .zip.",
             f"Deixe as duas tabelas marcadas, ligadas aos arquivos **{arq[T_SERV]}** e **{arq[T_VIG]}**.",
             f"Confira se ficaram **{n[T_SERV]} campos** em FOSERVICOS e **{n[T_VIG]} campos** em "
             "FOVIGENCIAS_SERVICO.",
         ],
         "dica": "Algum campo ficou sem tipo de dado? A importação parou no meio: apague os campos da tabela e "
                 "importe o arquivo de novo."},
        {"icone": "📥", "titulo": "Etapa 2 · Importar os tomadores",
         "quando": "Sempre que gerar um arquivo novo.",
         "passos": [
             f"Em **Arquivo > Serviços** da empresa **{emp}**, confira se {livre}.",
             f"Abra a importação por leiaute e escolha o leiaute **{lei}**.",
             f"Importe os arquivos da pasta **dados** do .zip, nesta ordem: **1º {arq[T_SERV]}** e "
             f"**2º {arq[T_VIG]}**.",
         ],
         "dica": "A ordem importa: a vigência só entra depois que o serviço já existe."},
        {"icone": "✅", "titulo": "Etapa 3 · Conferir na Domínio",
         "quando": "Logo depois da importação.",
         "passos": [
             "Em **Arquivo > Serviços**, abra um dos tomadores importados.",
             "Confira as guias **Geral** (CNPJ, endereço e CNAE), **eSocial** (tipo de serviço e lotação) e "
             "**INSS** (FPAS, terceiros, RAT, SEFIP e GPS).",
             "Vincule os colaboradores em **Processos > Outros > Alocação de Serviço**.",
         ],
         "dica": "O evento S-1020 vai para o eSocial quando o primeiro colaborador é vinculado ao serviço."},
    ]


def leia_me(info, perfil_txt):
    cods = info["codigos"]
    partes = [
        "IMPORTAÇÃO DE TOMADORES DE SERVIÇO — DOMÍNIO FOLHA",
        f"Gerado em {dt.datetime.now():%d/%m/%Y %H:%M} | Empresa {info.get('empresa')} | Leiaute {info['destino']}",
        f"{len(cods)} tomador(es), {_faixa(cods)}.",
        "",
        "DADOS HERDADOS DA PRESTADORA (iguais em todos os tomadores):",
        f"   {perfil_txt}",
    ]
    for e in passos_importacao(info):
        partes += ["", e["titulo"].upper(), f"   ({e['quando']})"]
        partes += [f"   {i}. {_sem_negrito(p)}" for i, p in enumerate(e["passos"], 1)]
        if e.get("dica"):
            partes.append(f"   Dica: {_sem_negrito(e['dica'])}")
    partes += ["", "DETALHES TÉCNICOS",
               "   Arquivos de dados: um tomador por linha, campos separados por TAB, sem cabeçalho, codificação "
               "cp1252, datas DD/MM/AAAA e decimais sem vírgula (20,000% = 20000).",
               "   A posição de cada campo nos arquivos está em estrutura_arquivos.csv."]
    return "\r\n".join(partes) + "\r\n"


def planilha_conferencia(abas: dict) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for nome, df in abas.items():
            df.to_excel(w, sheet_name=nome[:31], index=False)
            ws = w.sheets[nome[:31]]
            for c in ws[1]:
                c.font = Font(bold=True)
                c.alignment = Alignment(wrap_text=True)
    return buf.getvalue()
